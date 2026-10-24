import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"
HEADING_SIZE = Pt(16)
CONTENT_SIZE = Pt(14)
LINE_SPACING = 1.5

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level, space_before=16, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    h = doc.add_heading(text, level=level)
    h.alignment = align
    h.paragraph_format.space_before = Pt(space_before)
    h.paragraph_format.space_after = Pt(space_after)
    h.paragraph_format.line_spacing = 1.15
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = FONT_NAME
        r.font.size = HEADING_SIZE # Exact 16 pt as requested
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42) # Professional Deep Slate / Black
    return h

def add_body_p(doc, text="", bold_prefix="", space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align # Exact Justify as requested
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = LINE_SPACING # Exact 1.5 line spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = FONT_NAME
        r_b.font.size = CONTENT_SIZE # Exact 14 pt as requested
        r_b.font.color.rgb = RGBColor(15, 23, 42)
    if text:
        r_t = p.add_run(text)
        r_t.font.name = FONT_NAME
        r_t.font.size = CONTENT_SIZE # Exact 14 pt as requested
        r_t.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, text="", bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY # Justified alignment
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = LINE_SPACING # 1.5 line spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = FONT_NAME
        r_b.font.size = CONTENT_SIZE # 14 pt
        r_b.font.color.rgb = RGBColor(15, 23, 42)
    if text:
        r_t = p.add_run(text)
        r_t.font.name = FONT_NAME
        r_t.font.size = CONTENT_SIZE # 14 pt
        r_t.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_image_box(doc, img_path, caption, width=Inches(6.2)):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found.")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(6)
    run = p_img.add_run()
    run.add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(14)
    p_cap.paragraph_format.line_spacing = 1.15
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = FONT_NAME
    r_cap.font.size = Pt(12)
    r_cap.font.italic = True
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(51, 65, 85)

def create_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = FONT_NAME
            r.font.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(11.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return tbl
